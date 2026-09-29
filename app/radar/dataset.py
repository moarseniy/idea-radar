from __future__ import annotations

import csv
import math
import re
from collections import Counter
from functools import lru_cache
from pathlib import Path

from app.ml.features.metadata import FEATURE_DESCRIPTIONS, FEATURE_GROUPS, FEATURE_GROUP_BY_NAME
from app.ml.features.schema import FEATURE_SCHEMA_VERSION
from app.ml.scorer import FEATURE_LABELS, default_scorer, display_value


@lru_cache(maxsize=4)
def _read_features(path: str, modified: int, feature_names: tuple[str, ...]) -> dict[str, list[dict]]:
    result = {}
    with Path(path).open(newline="", encoding="utf-8-sig") as stream:
        for row in csv.DictReader(stream):
            if row.get("feature_schema_version") != FEATURE_SCHEMA_VERSION:
                continue
            features = []
            for name in feature_names:
                raw = row.get(name, "").strip()
                try:
                    value = float(raw) if raw else None
                    if value is not None and not math.isfinite(value):
                        value = None
                except ValueError:
                    value = None
                features.append({
                    "name": name,
                    "label": FEATURE_LABELS[name],
                    "description": FEATURE_DESCRIPTIONS[name],
                    "group": FEATURE_GROUP_BY_NAME[name],
                    "group_label": FEATURE_GROUPS[FEATURE_GROUP_BY_NAME[name]],
                    "value": value,
                    "display_value": display_value(name, value),
                })
            result[row.get("organizer_id", "")] = features
    return result


def _with_features(records: list[dict], data_dir: Path, feature_names: tuple[str, ...]) -> list[dict]:
    path = data_dir / "training" / "organizer_100_features.csv"
    feature_rows = (_read_features(str(path.resolve()), path.stat().st_mtime_ns, feature_names)
                    if path.is_file() else {})
    for record in records:
        features = feature_rows.get(f"organizer_{record['id']:03d}", [])
        record["features"] = features
        record["measured_feature_count"] = sum(feature["value"] is not None for feature in features)
    return records


@lru_cache(maxsize=4)
def _read(path: str, modified: int) -> dict:
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=True)
    records = []
    try:
        for sheet in workbook:
            header = None
            for row_no, row in enumerate(sheet.iter_rows(values_only=True), 1):
                values = [str(v).strip() if v is not None else "" for v in row]
                if "Технология (слабый сигнал)" in values and "Источники" in values:
                    header = {name: i for i, name in enumerate(values) if name}
                    continue
                if not header:
                    continue

                def cell(name):
                    index = header.get(name)
                    return row[index] if index is not None and index < len(row) else None

                if not isinstance(cell("№"), (int, float)):
                    continue
                links = [{"title": title, "url": url} for title, url in re.findall(
                    r"\[([^\]]+)\]\((https?://[^\s)]+)\)", str(cell("Источники") or ""))]
                records.append({"id": int(cell("№")), "title": cell("Технология (слабый сигнал)"),
                    "domain": cell("Область"), "companies": cell("Компании"),
                    "reference_explanation": cell("Почему это слабый сигнал"), "stage": cell("Стадия развития"),
                    "mention_trend": cell("Тренд упоминаний"), "reference_score": cell("Балл (стадия+тренд)"),
                    "sources": links, "sheet": sheet.title, "row": row_no,
                    "origin": "organizer_catalog", "independently_verified": False})
    finally:
        workbook.close()
    return {"filename": Path(path).name, "records": records, "count": len(records),
            "domains": dict(Counter(r["domain"] for r in records)),
            "note": "Примеры организаторов. Объяснения и баллы сохранены из Excel; это не результаты живого поиска и не независимая валидация."}


def read_catalog(data_dir: Path) -> dict:
    path = next(iter(sorted(data_dir.glob("*сигнал*.xlsx"))), None)
    if path is None:
        return {"records": [], "count": 0, "domains": {}, "note": "Файл примеров не найден."}
    catalog = _read(str(path.resolve()), path.stat().st_mtime_ns)
    feature_names = tuple(default_scorer().features)
    records = _with_features([dict(record) for record in catalog["records"]], data_dir, feature_names)
    return {**catalog, "records": records,
            "feature_count": len(feature_names),
            "feature_groups": FEATURE_GROUPS}
